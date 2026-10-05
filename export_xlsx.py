"""
Exporte fut27.db vers un fichier .xlsx avec exactement les mêmes colonnes que
le pilote manuel Futbin (Nom prenom, Note, Pays, Ligue, Club, PS+, PS,
Positions, les 6 stats principales en français, puis toutes les sous-stats
détaillées) - mais pour autant de cartes que tu veux, sans copier-coller
Futbin à la main.

Nécessite d'avoir lancé, dans l'ordre :
    python scrape_players.py --out fut27.db
    python scrape_player_details.py --db fut27.db   (au moins pour les cartes voulues)

Usage:
    python export_xlsx.py --db fut27.db --out export.xlsx
    python export_xlsx.py --db fut27.db --out export.xlsx --min-overall 75 --max-overall 99
    python export_xlsx.py --db fut27.db --out export.xlsx --position ST
"""
import argparse
import json
import sqlite3

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from evolution_report import parse_alt_positions

MAIN_CATS = ["Pace", "Shooting", "Passing", "Dribbling", "Defending", "Physical"]
MAIN_LABELS_FR = {
    "Pace": "Vitesse",
    "Shooting": "Tir",
    "Passing": "Passe",
    "Dribbling": "Dribbles",
    "Defending": "Defense",
    "Physical": "Physique",
}


def load_rows(conn: sqlite3.Connection, min_overall: int, max_overall: int, position: str | None) -> list[dict]:
    query = """
        SELECT p.common_name, p.overall, p.nation, p.league, p.club,
               p.position, p.alternative_positions,
               p.playstyles_plus, p.playstyles_normal,
               pa.detail_stats_json
        FROM players p
        JOIN player_attributes pa ON pa.player_id = p.id
        WHERE p.overall BETWEEN ? AND ?
    """
    params: list = [min_overall, max_overall]
    if position:
        query += " AND p.position = ?"
        params.append(position)
    query += " ORDER BY p.overall DESC"

    out = []
    for row in conn.execute(query, params):
        (
            name, overall, nation, league, club,
            position1, alt_positions_raw, ps_plus_raw, ps_normal_raw, detail_json,
        ) = row
        alt_positions = parse_alt_positions(alt_positions_raw)
        out.append(
            {
                "name": name,
                "rating": overall,
                "nation": nation,
                "league": league,
                "club": club,
                "position1": position1,
                "altPositions": alt_positions,
                "psPlus": json.loads(ps_plus_raw or "[]"),
                "psNormal": json.loads(ps_normal_raw or "[]"),
                "statPairs": json.loads(detail_json or "[]"),
            }
        )
    return out


def build_workbook(players: list[dict], out_path: str) -> int:
    max_alt = max((len(p["altPositions"]) for p in players), default=0)
    max_psplus = max((len(p["psPlus"]) for p in players), default=0)
    max_psnormal = max((len(p["psNormal"]) for p in players), default=0)

    detail_columns: list[str] = []
    for p in players:
        seen_main = set()
        for label, _ in p["statPairs"]:
            if label in MAIN_CATS and label not in seen_main:
                seen_main.add(label)
                continue
            if label not in detail_columns:
                detail_columns.append(label)

    columns = (
        ["Nom prenom", "Note", "Pays", "Ligue", "Club"]
        + [f"PS+{i + 1}" for i in range(max_psplus)]
        + [f"PS{i + 1}" for i in range(max_psnormal)]
        + [f"Position{i + 1}" for i in range(1 + max_alt)]
        + [MAIN_LABELS_FR[c] for c in MAIN_CATS]
        + detail_columns
    )

    rows = []
    for p in players:
        seen_main = set()
        detail_values: dict = {}
        main_values: dict = {}
        for label, val in p["statPairs"]:
            if label in MAIN_CATS and label not in seen_main:
                seen_main.add(label)
                main_values[label] = val
                continue
            detail_values[label] = val

        positions = [p["position1"]] + p["altPositions"]
        positions += [""] * (1 + max_alt - len(positions))

        psplus_cols = p["psPlus"] + [""] * (max_psplus - len(p["psPlus"]))
        psnormal_cols = p["psNormal"] + [""] * (max_psnormal - len(p["psNormal"]))

        row = (
            [p["name"], p["rating"], p["nation"], p["league"], p["club"]]
            + psplus_cols
            + psnormal_cols
            + positions
            + [main_values.get(c, "") for c in MAIN_CATS]
            + [detail_values.get(c, "") for c in detail_columns]
        )
        rows.append(row)

    wb = Workbook()
    ws = wb.active
    ws.title = "Joueurs"
    ws.append(columns)
    for row in rows:
        ws.append(row)

    for i, col in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(10, min(22, len(col) + 2))
    ws.freeze_panes = "A2"

    wb.save(out_path)
    return len(rows)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--db", default="fut27.db")
    parser.add_argument("--out", default="export.xlsx")
    parser.add_argument("--min-overall", type=int, default=1)
    parser.add_argument("--max-overall", type=int, default=99)
    parser.add_argument("--position", default=None, help="ex: ST, CB, GK...")
    args = parser.parse_args()

    conn = sqlite3.connect(args.db)
    players = load_rows(conn, args.min_overall, args.max_overall, args.position)
    if not players:
        print("Aucune carte trouvee avec ces filtres (as-tu lance scrape_player_details.py ?)")
        return

    n = build_workbook(players, args.out)
    print(f"Ecrit {n} lignes dans {args.out}")


if __name__ == "__main__":
    main()
